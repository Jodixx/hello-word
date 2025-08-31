import json
from pathlib import Path
import argparse
from typing import List, Dict, Any

try:
    from fpdf import FPDF
except ImportError:  # pragma: no cover - handled during runtime
    FPDF = None

DATA_FILE = Path("session_data.json")

MODULE_NAMES = [
    "Definição do Negócio",
    "Público-Alvo",
    "Necessidades do Público",
    "Atributos Funcionais",
    "Atributos Emocionais",
    "Valores",
    "Canais Estratégicos",
    "Propósito e Análise Competitiva",
]

MODULE_KEYS = [
    "business_definition",
    "target_audience",
    "audience_needs",
    "functional_attributes",
    "emotional_attributes",
    "values",
    "strategic_channels",
    "purpose_competition",
]

def save_data(data: Dict[str, Any]) -> None:
    with DATA_FILE.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def load_data() -> Dict[str, Any]:
    if DATA_FILE.exists():
        with DATA_FILE.open("r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def show_progress(step: int, total: int) -> None:
    bar_len = 30
    filled = int(bar_len * step / total)
    bar = "=" * filled + " " * (bar_len - filled)
    print(f"[{bar}] {step}/{total}")


def show_checklist(data: Dict[str, Any]) -> None:
    for idx, name in enumerate(MODULE_NAMES):
        key = MODULE_KEYS[idx]
        status = "✓" if key in data else " "
        print(f"[{status}] {idx + 1}. {name}")


def simulate_trends(phrase: str) -> List[str]:
    words = phrase.split()[:3]
    return [f"{w.lower()} tendência" for w in words]


def module1(data: Dict[str, Any]) -> None:
    print("\nMódulo 1 – Definição do Negócio")
    print("Estrutura sugerida: 'O que + Para quem + Como'")
    phrase = input("Frase curta que define o negócio: ").strip()
    keywords = simulate_trends(phrase)
    print("Palavras-chave sugeridas:", ", ".join(keywords))
    data[MODULE_KEYS[0]] = {"frase": phrase, "palavras_chave": keywords}


def simulate_audience(age, gender, location, income, interests, values):
    # Placeholder for audience tool
    size = 1000 + len(interests.split(",")) * 100
    trends = [i.strip().title() for i in interests.split(",") if i.strip()]
    return {"tamanho": size, "tendencias": trends, "gaps": "N/A"}


def module2(data: Dict[str, Any]) -> None:
    print("\nMódulo 2 – Público-Alvo")
    age = input("Idade: ")
    gender = input("Gênero: ")
    location = input("Localização: ")
    income = input("Faixa de renda: ")
    interests = input("Principais interesses (separados por vírgula): ")
    values = input("Valores principais (separados por vírgula): ")
    audience_data = simulate_audience(age, gender, location, income, interests, values)
    print("Análise simulada de público:")
    print(audience_data)
    data[MODULE_KEYS[1]] = {
        "idade": age,
        "genero": gender,
        "localizacao": location,
        "renda": income,
        "interesses": interests,
        "valores": values,
        "analise": audience_data,
    }


def simulate_need_validation(need: str) -> str:
    return f"A dor '{need}' possui relevância moderada."  # placeholder


def module3(data: Dict[str, Any]) -> None:
    print("\nMódulo 3 – Necessidades do Público")
    needs: Dict[str, str] = {}
    for i in range(5):
        need = input(f"Dor/Problema #{i + 1} (deixe vazio para parar): ").strip()
        if not need:
            break
        solution = input("Como a marca resolve isso? ")
        validation = simulate_need_validation(need)
        print(validation)
        needs[need] = {"solucao": solution, "validacao": validation}
    data[MODULE_KEYS[2]] = needs


def simulate_competitive(benefit: str) -> str:
    return f"'{benefit}' é pouco explorado." if len(benefit) % 2 else f"'{benefit}' é comum."  # placeholder


def module4(data: Dict[str, Any]) -> None:
    print("\nMódulo 4 – Atributos Funcionais")
    benefits = []
    for i in range(5):
        b = input(f"Benefício prático #{i + 1} (deixe vazio para parar): ").strip()
        if not b:
            break
        analysis = simulate_competitive(b)
        print(analysis)
        benefits.append({"beneficio": b, "analise": analysis})
    data[MODULE_KEYS[3]] = benefits


def simulate_emotional_alignment(emotion: str) -> str:
    return f"Emoção '{emotion}' está alinhada."  # placeholder


def module5(data: Dict[str, Any]) -> None:
    print("\nMódulo 5 – Atributos Emocionais")
    emotions = []
    for i in range(5):
        em = input(f"Emoção #{i + 1} (deixe vazio para parar): ").strip()
        if not em:
            break
        justification = input("Breve justificativa/ação: ")
        analysis = simulate_emotional_alignment(em)
        print(analysis)
        emotions.append({"emocao": em, "justificativa": justification, "analise": analysis})
    data[MODULE_KEYS[4]] = emotions


def simulate_value_relevance(value: str) -> str:
    return f"Valor '{value}' é relevante."  # placeholder


def module6(data: Dict[str, Any]) -> None:
    print("\nMódulo 6 – Valores")
    values_list = []
    for i in range(5):
        val = input(f"Valor #{i + 1} (deixe vazio para parar): ").strip()
        if not val:
            break
        example = input("Exemplo prático: ")
        analysis = simulate_value_relevance(val)
        print(analysis)
        values_list.append({"valor": val, "exemplo": example, "analise": analysis})
    data[MODULE_KEYS[5]] = values_list


def simulate_channel_priority(channels: List[str], budget: str, objective: str) -> str:
    return f"Priorize: {', '.join(channels[:2])}. Orçamento: {budget}. Objetivo: {objective}."


def module7(data: Dict[str, Any]) -> None:
    print("\nMódulo 7 – Canais Estratégicos")
    channels = input("Canais selecionados (separados por vírgula): ").split(",")
    budget = input("Orçamento disponível: ")
    objective = input("Objetivo principal: ")
    suggestion = simulate_channel_priority(channels, budget, objective)
    print("Sugestão automática:", suggestion)
    data[MODULE_KEYS[6]] = {
        "canais": [c.strip() for c in channels if c.strip()],
        "orcamento": budget,
        "objetivo": objective,
        "sugestao": suggestion,
    }


def simulate_competitors(purpose: str) -> Dict[str, Any]:
    competitors = [
        {"nome": "Concorrente A", "forcas": "Preço", "fraquezas": "Atendimento"},
        {"nome": "Concorrente B", "forcas": "Marca", "fraquezas": "Preço"},
    ]
    return {"proposito": purpose, "concorrentes": competitors, "resumo": "Resumo executivo"}


def module8(data: Dict[str, Any]) -> None:
    print("\nMódulo 8 – Propósito e Análise Competitiva")
    purpose = input("Frase de propósito: ")
    promise = input("Promessa da marca: ")
    analysis = simulate_competitors(purpose)
    print("Análise de concorrentes simulada:")
    for c in analysis["concorrentes"]:
        print(c)
    data[MODULE_KEYS[7]] = {"proposito": purpose, "promessa": promise, "analise": analysis}


def generate_pdf(data: Dict[str, Any], filename: str = "estrategia.pdf") -> None:
    if FPDF is None:
        print("Biblioteca FPDF não instalada; PDF não gerado.")
        return
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "Estratégia de Marca", ln=True)
    pdf.set_font("Arial", size=12)
    for key, value in data.items():
        pdf.cell(0, 10, key.replace('_', ' ').title(), ln=True)
        pdf.set_font_size(10)
        pdf.multi_cell(0, 5, json.dumps(value, ensure_ascii=False, indent=2))
        pdf.set_font_size(12)
    pdf.output(filename)
    print(f"PDF gerado: {filename}")


def export_powerpoint(data: Dict[str, Any], filename: str = "resumo.pptx") -> None:
    try:
        from pptx import Presentation
    except Exception:  # pragma: no cover - optional dependency
        print("Dependência 'python-pptx' não instalada; PPT não gerado.")
        return
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Resumo de Estratégia"
    body = slide.placeholders[1].text_frame
    for key in MODULE_KEYS:
        if key in data:
            body.add_paragraph().text = f"{key}: {str(data[key])[:30]}"
    prs.save(filename)
    print(f"PowerPoint gerado: {filename}")


def run_interactive() -> None:
    data = load_data()
    segment = data.get("segmento") or input("Segmento de mercado: ")
    data["segmento"] = segment
    step = 0
    total = len(MODULE_NAMES)
    modules = [module1, module2, module3, module4, module5, module6, module7, module8]
    while step < total:
        show_checklist(data)
        show_progress(step, total)
        modules[step](data)
        save_data(data)
        action = input("Digite 'voltar' para etapa anterior, 'editar' para refazer ou Enter para continuar: ").strip().lower()
        if action == "voltar" and step > 0:
            step -= 1
        elif action == "editar":
            continue
        else:
            step += 1
    generate_pdf(data)
    if input("Exportar para PowerPoint? (s/n): ").strip().lower() == 's':
        export_powerpoint(data)
    print("Processo concluído!")


def run_demo() -> None:
    demo_data = {
        "segmento": "Exemplo",
        MODULE_KEYS[0]: {"frase": "Loja online de roupas", "palavras_chave": ["loja", "online"]},
        MODULE_KEYS[1]: {"idade": "25-40", "genero": "Todos", "analise": {"tamanho": 2000}},
        MODULE_KEYS[2]: {"Entrega rápida": {"solucao": "Entrega em 24h"}},
        MODULE_KEYS[3]: [{"beneficio": "Frete grátis"}],
        MODULE_KEYS[4]: [{"emocao": "Confiança"}],
        MODULE_KEYS[5]: [{"valor": "Transparência"}],
        MODULE_KEYS[6]: {"canais": ["Instagram", "Google Ads"]},
        MODULE_KEYS[7]: {"proposito": "Vestir bem", "promessa": "Qualidade"},
    }
    generate_pdf(demo_data, filename="demo_estrategia.pdf")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Ferramenta de Estratégia de Marca")
    parser.add_argument("--demo", action="store_true", help="Executa geração de relatório com dados de demonstração")
    args = parser.parse_args()
    if args.demo:
        run_demo()
    else:
        run_interactive()
